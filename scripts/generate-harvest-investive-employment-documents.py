from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


DESKTOP = Path.home() / "Desktop"
OUTPUT_DIR = DESKTOP / "Harvest Investive Employment Documents"
REPORTS_DIR = Path("reports") / "Harvest Investive Employment Documents"
TODAY = date.today().strftime("%d %B %Y")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "D0D7DE")


def setup_doc(title, subtitle=None):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.1

    for style_name, size, color, before, after in [
        ("Heading 1", 15, "1F4D78", 12, 6),
        ("Heading 2", 12.5, "2E74B5", 9, 4),
        ("Heading 3", 11.5, "1F4D78", 6, 3),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor(11, 37, 69)

    if subtitle:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(subtitle)
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(80, 86, 94)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"Prepared for Harvest Investive | Draft date: {TODAY}")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(100, 108, 118)
    return doc


def add_note(doc):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F4F6F9")
    p = cell.paragraphs[0]
    r = p.add_run(
        "Legal review note: This is a practical draft for internal use. Before signing, Harvest Investive should confirm the correct NEC/CBA, minimum wage, leave, notice, and disciplinary-code requirements applicable to the business and location."
    )
    r.bold = True
    r.font.size = Pt(9)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_metadata_table(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    for label, value in rows:
        cells = table.add_row().cells
        set_cell_width(cells[0], 2400)
        set_cell_width(cells[1], 6960)
        cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cells[0], "F2F4F7")
        p = cells[0].paragraphs[0]
        p.add_run(label).bold = True
        cells[1].paragraphs[0].add_run(value)
    return table


def add_signature_block(doc, employee_name, company_signatory="David Chaeruka / Authorised Representative"):
    doc.add_heading("Signatures", level=1)
    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    labels = [
        ("For Harvest Investive", "Employee"),
        (company_signatory, employee_name),
        ("Signature: ______________________________", "Signature: ______________________________"),
        ("Date: __________________", "Date: __________________"),
    ]
    for row, values in zip(table.rows, labels):
        for idx, text in enumerate(values):
            set_cell_width(row.cells[idx], 4680)
            row.cells[idx].paragraphs[0].add_run(text)


def add_contract_common(doc, employee):
    doc.add_heading("1. Appointment", level=1)
    doc.add_paragraph(
        f"Harvest Investive appoints {employee['name']} as {employee['role']} assigned to {employee['shop']} on the terms set out in this contract. The employee accepts the appointment and agrees to perform duties honestly, carefully, and in the best interests of Harvest Investive."
    )

    doc.add_heading("2. Key Employment Particulars", level=1)
    add_metadata_table(
        doc,
        [
            ("Employer", "Harvest Investive"),
            ("Employee", employee["name"]),
            ("Role", employee["role"]),
            ("Assigned shop/base", employee["shop"]),
            ("Start date", "[Insert start date]"),
            ("Employment type", employee.get("employment_type", "Full-time permanent, subject to probation")),
            ("Probation", "Three months, subject to lawful review and fair process"),
            ("Basic salary/wage", "[Insert amount and pay frequency]"),
            ("Normal working days/hours", "[Insert roster, opening hours, lunch/rest arrangements, and weekly rest day]"),
            ("Reporting line", employee.get("reporting", "Managing Director / Shop Supervisor")),
        ],
    )

    doc.add_heading("3. Duties and Performance Standards", level=1)
    add_bullets(doc, employee["duties"])
    add_bullets(
        doc,
        [
            "Follow all Harvest Investive policies, shop rules, stock controls, POS controls, cash controls, and lawful instructions.",
            "Protect company stock, cash, equipment, records, passwords, supplier information, and customer information.",
            "Report suspected theft, cash variance, stock loss, fraud, misconduct, or unsafe work conditions immediately.",
        ],
    )

    doc.add_heading("4. Cash, Stock, and POS Accountability", level=1)
    add_bullets(
        doc,
        [
            "No employee may remove cash, stock, receipts, devices, or records without written approval.",
            "All sales must be entered into the approved POS or written manual backup register before stock leaves the shop.",
            "Cash drawer counts, handovers, banking, savings, black box, rent reserve, and Operations transfers must match the daily record.",
            "Unexplained shortages, false entries, deliberate under-ringing, unauthorised discounts, or private use of company money may lead to disciplinary action.",
        ],
    )

    doc.add_heading("5. Pay, Benefits, Deductions, and Records", level=1)
    doc.add_paragraph(
        "Salary, allowances, benefits, commission, deductions, and payment dates must be recorded in writing. Any lawful deductions must be explained to the employee and supported by company records, applicable law, or written agreement where required."
    )

    doc.add_heading("6. Leave, Public Holidays, Sickness, and Absence", level=1)
    add_bullets(
        doc,
        [
            "Leave, public holidays, sick leave, and other statutory benefits will be handled in line with applicable Zimbabwean law, the relevant NEC/CBA where applicable, and company policy.",
            "Absence must be reported to the reporting line before shift start unless impossible.",
            "Sick leave may require medical proof where lawful and reasonable.",
        ],
    )

    doc.add_heading("7. Confidentiality and Conflict of Interest", level=1)
    add_bullets(
        doc,
        [
            "The employee must keep business information confidential during and after employment.",
            "The employee must not run competing sales, divert customers, use company suppliers privately, or take side payments from suppliers or customers.",
            "Any family, supplier, customer, or personal relationship that may affect business decisions must be disclosed.",
        ],
    )

    doc.add_heading("8. Discipline, Grievance, and Termination", level=1)
    add_bullets(
        doc,
        [
            "Discipline and grievance matters will be handled fairly and in line with applicable law, any registered employment code, and the relevant NEC/CBA where applicable.",
            "Termination may occur through resignation, mutual agreement, expiry of a fixed term if applicable, dismissal after fair process, retrenchment, or any other lawful ground.",
            "On termination, the employee must return all company property, cash, stock, records, uniforms, keys, passwords, and devices immediately.",
        ],
    )

    doc.add_heading("9. Entire Agreement", level=1)
    doc.add_paragraph(
        "This contract, the Harvest Investive Work Policy, and any written salary or role schedule form the employment agreement. Any changes must be recorded in writing and signed or acknowledged by the parties."
    )


def create_contract(employee):
    title = "Employment Contract"
    subtitle = f"{employee['name']} - {employee['role']} - {employee['shop']}"
    doc = setup_doc(title, subtitle)
    add_note(doc)
    add_contract_common(doc, employee)
    add_signature_block(doc, employee["name"], employee.get("company_signatory", "Authorised Representative"))
    return doc


def create_policy():
    doc = setup_doc("Harvest Investive Work Policy", "Applies to Dub Dub, Kipasa, and Trade Center")
    add_note(doc)

    doc.add_heading("1. Purpose", level=1)
    doc.add_paragraph(
        "This policy creates one operating standard for Harvest Investive shops: Dub Dub, Kipasa, and Trade Center. It is intended to protect cash, stock, staff, customers, records, and owner capital."
    )

    doc.add_heading("2. Core Rules", level=1)
    add_bullets(
        doc,
        [
            "Every sale must be recorded before stock leaves the shop.",
            "Every cash movement must have a reason, amount, person responsible, date, and proof.",
            "No Operations, savings, black box, rent reserve, stock purchase, or banking entry may be made without a matching note or receipt.",
            "No one may borrow company cash or stock without written owner approval.",
            "Each shop must close daily with a signed cash count, stock issue note where relevant, and variance explanation.",
        ],
    )

    doc.add_heading("3. Shop Roles and Accountability", level=1)
    add_metadata_table(
        doc,
        [
            ("Dub Dub", "Two assigned shop employees. Responsible for POS, customer service, cash drawer, stock display, and daily handover."),
            ("Trade Center", "Two assigned shop employees. Responsible for POS, customer service, cash drawer, stock display, and daily handover."),
            ("Kipasa", "Two assigned shop employees. Responsible for POS, customer service, cash drawer, stock display, and daily handover."),
            ("MD/CEO", "Responsible for financial discipline, owner reporting, approvals, staff oversight, procurement discipline, and policy enforcement."),
        ],
    )

    doc.add_heading("4. Cash Handling Policy", level=1)
    add_numbered(
        doc,
        [
            "Each shift starts with an opening cash count signed by the person taking responsibility for the drawer.",
            "All cash sales, EcoCash/mobile money sales, lay-bys, refunds, expenses, savings, black box, rent reserve, and Operations transfers must be recorded the same day.",
            "Cash may not be mixed with personal money.",
            "Cash shortages must be reported immediately and recorded before the shop closes.",
            "Cash overages are company money and must be recorded. Staff may not keep overages.",
            "Banking must be supported by a deposit slip, bank confirmation, or owner acknowledgement.",
        ],
    )

    doc.add_heading("5. Operations, Savings, Black Box, and Rent Reserve", level=1)
    add_bullets(
        doc,
        [
            "Savings and black box money must be treated as protected business money, not casual spend money.",
            "Rent reserve must be ring-fenced monthly. If the company budgets USD 2,000 and pays USD 1,200 rent, the USD 800 difference must remain visible in a reserve record.",
            "Operations payouts require a category: stock order, rent, tax, banking, salaries, utilities, transport, supplier/payable, maintenance, or owner-approved other.",
            "Generic labels such as 'Round', 'Balance', 'Person name only', 'Adjustment', or 'Other' are not acceptable without supporting explanation.",
            "Any Operations payout over USD 100 requires proof. Any payout over USD 500 requires owner approval and proof before or immediately after payment.",
        ],
    )

    doc.add_heading("6. Stock Purchasing and Inventory", level=1)
    add_bullets(
        doc,
        [
            "Stock orders must have supplier, amount, date, receiving shop, expected items, and proof of payment.",
            "Stock received must be checked against the order before it is sold.",
            "Damaged, missing, returned, or transferred stock must be recorded.",
            "Staff may not take stock on credit or for personal use unless authorised in writing.",
        ],
    )

    doc.add_heading("7. Expenses, Salaries, and Approvals", level=1)
    add_bullets(
        doc,
        [
            "Routine shop expenses must be pre-approved where possible and recorded with receipts.",
            "Salary and wage payments must be recorded separately from stock, rent, banking, and owner drawings.",
            "Tax must be recorded separately from banking. The business must not hide ordinary cash movement inside a tax/banking label.",
            "Owner drawings, loans, or advances must be labelled separately and cannot be treated as operating expenses.",
        ],
    )

    doc.add_heading("8. Attendance, Conduct, and Customer Service", level=1)
    add_bullets(
        doc,
        [
            "Staff must report on time, dress neatly, serve customers respectfully, and keep the shop clean.",
            "Staff must not drink alcohol, use drugs, gamble, fight, harass others, or bring the company into disrepute during work.",
            "Phones may be used for business communication but must not interfere with customers, cash handling, or stock control.",
            "Customer complaints must be escalated to the MD/CEO or owner.",
        ],
    )

    doc.add_heading("9. Discipline and Grievance Procedure", level=1)
    add_bullets(
        doc,
        [
            "Misconduct will be handled fairly and proportionately.",
            "Serious misconduct may include theft, fraud, falsifying records, unauthorised cash use, violence, gross negligence, or deliberate stock loss.",
            "Employees may raise grievances without victimisation. The complaint must be recorded, investigated, and responded to.",
            "Where an NEC/CBA or registered employment code applies, it must be followed.",
        ],
    )

    doc.add_heading("10. Daily Required Records", level=1)
    add_bullets(
        doc,
        [
            "Opening cash count.",
            "Daily sales report.",
            "Expense and payout register with receipts.",
            "Savings, black box, rent reserve, and Operations transfer register.",
            "Closing cash count and variance note.",
            "Stock received, stock transferred, and stock lost/damaged register.",
        ],
    )

    doc.add_heading("11. Acknowledgement", level=1)
    doc.add_paragraph(
        "Each employee must sign acknowledgement that they received, understood, and agreed to follow this policy. Refusal to follow the policy may lead to disciplinary action."
    )
    add_signature_block(doc, "Employee", "David Chaeruka / Authorised Representative")
    return doc


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)

    employees = [
        {
            "name": "David Chaeruka",
            "role": "Managing Director / Chief Executive Officer",
            "shop": "Harvest Investive Head Office / All Shops",
            "reporting": "Owner / Shareholders",
            "company_signatory": "Owner / Authorised Representative",
            "duties": [
                "Lead daily operations across Dub Dub, Kipasa, and Trade Center.",
                "Enforce cash controls, rent reserve controls, stock purchasing controls, and staff discipline.",
                "Prepare weekly owner reports covering sales, cash, stock, expenses, rent reserve, black box, savings, and Operations payouts.",
                "Approve staff rosters, investigate variances, maintain supplier discipline, and escalate risks promptly.",
                "Protect company strategy, finances, staff records, supplier relationships, and owner instructions.",
            ],
        },
        {
            "name": "[Dub Dub Employee 1]",
            "role": "Shop Assistant / Cashier",
            "shop": "Dub Dub",
            "duties": [
                "Open and close the shop as rostered.",
                "Record sales, receive customers, maintain displays, and protect the cash drawer.",
                "Record savings, black box, rent reserve, and Operations transfers exactly as instructed.",
                "Assist with stock counts, stock receiving, cleaning, and daily handover.",
            ],
        },
        {
            "name": "[Dub Dub Employee 2]",
            "role": "Shop Assistant / Cashier",
            "shop": "Dub Dub",
            "duties": [
                "Serve customers, process sales, maintain shop cleanliness, and protect stock.",
                "Support cash counts, POS entries, expense records, and handover notes.",
                "Report all variances, customer complaints, stock issues, and suspicious activity immediately.",
            ],
        },
        {
            "name": "[Trade Center Employee 1]",
            "role": "Shop Assistant / Cashier",
            "shop": "Trade Center",
            "duties": [
                "Open and close the Trade Center shop as rostered.",
                "Record sales accurately and maintain customer service standards.",
                "Handle cash, stock, and transfer records according to policy.",
                "Assist with stock checks, receiving, cleaning, and daily reporting.",
            ],
        },
        {
            "name": "[Trade Center Employee 2]",
            "role": "Shop Assistant / Cashier",
            "shop": "Trade Center",
            "duties": [
                "Serve customers, record sales, protect stock, and support shop presentation.",
                "Complete cash and stock handovers truthfully.",
                "Report shortages, overages, refunds, returns, and damaged stock immediately.",
            ],
        },
        {
            "name": "[Kipasa Employee 1]",
            "role": "Shop Assistant / Cashier",
            "shop": "Kipasa",
            "duties": [
                "Open and close the Kipasa shop as rostered.",
                "Record sales, maintain displays, protect cash, and manage customer service.",
                "Support daily cash count, stock count, transfer records, and expense documentation.",
            ],
        },
        {
            "name": "[Kipasa Employee 2]",
            "role": "Shop Assistant / Cashier",
            "shop": "Kipasa",
            "duties": [
                "Serve customers and record all sales before stock leaves the shop.",
                "Protect the cash drawer and company stock during shifts.",
                "Report stock issues, cash variances, and suspicious activity immediately.",
            ],
        },
    ]

    created = []
    docs = [("Harvest Investive - Work Policy - Dub Dub Kipasa Trade Center.docx", create_policy())]
    for employee in employees:
        safe_name = employee["name"].replace("[", "").replace("]", "").replace("/", "-")
        docs.append((f"Employment Contract - {safe_name} - {employee['shop'].replace('/', '-')}.docx", create_contract(employee)))

    for filename, doc in docs:
        desktop_path = OUTPUT_DIR / filename
        report_path = REPORTS_DIR / filename
        doc.save(desktop_path)
        doc.save(report_path)
        created.append(str(desktop_path))

    manifest = OUTPUT_DIR / "README - Document Set.txt"
    manifest.write_text(
        "Harvest Investive employment document set\\n"
        f"Generated: {TODAY}\\n\\n"
        "Files included:\\n"
        + "\\n".join(f"- {Path(item).name}" for item in created)
        + "\\n\\nBefore signing: fill employee legal names, ID numbers, addresses, salary/wage, start date, working hours, leave details, and signatures. Confirm NEC/CBA requirements and local legal compliance.\\n",
        encoding="utf-8",
    )

    print("\\n".join(created))


if __name__ == "__main__":
    main()
