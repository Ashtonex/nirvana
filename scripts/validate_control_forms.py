from pathlib import Path
from zipfile import ZipFile

from docx import Document


def main():
    for path in sorted(Path("docs/control-forms").glob("*.docx")):
        with ZipFile(path) as archive:
            bad = archive.testzip()
            names = set(archive.namelist())
        doc = Document(path)
        paragraphs = sum(1 for p in doc.paragraphs if p.text.strip())
        print(
            f"{path}: "
            f"zip_ok={bad is None}, "
            f"document_xml={'word/document.xml' in names}, "
            f"tables={len(doc.tables)}, "
            f"paragraphs={paragraphs}, "
            f"sections={len(doc.sections)}, "
            f"size={path.stat().st_size}"
        )


if __name__ == "__main__":
    main()
