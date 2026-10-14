from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT_DIR = Path("docs/control-forms")
BRAND = "INVENTIVE HARVEST DUB-KIPASA-TC"

BLUE = "1F4D78"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
BORDER = "8EA9C1"
INK = "0B2545"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER, size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for w in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(w))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.tcW
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            set_cell_border(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def style_paragraph(p, size=9.5, bold=False, color="000000", align=None, before=0, after=2):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.05
    if align is not None:
        p.alignment = align
    for run in p.runs:
        run.font.name = "Calibri"
        run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)


def cell_text(cell, text, size=9, bold=False, color="000000", align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    style_paragraph(p, size=size, bold=bold, color=color, align=align, after=0)


def add_rule_paragraph(doc, label, length=44):
    p = doc.add_paragraph()
    run = p.add_run(f"{label}: " + "_" * length)
    style_paragraph(p, size=9.5, after=4)
    return p


def setup_doc(title, subtitle):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    for attr in ("top_margin", "right_margin", "bottom_margin", "left_margin"):
        setattr(sec, attr, Inches(0.72))
    sec.header_distance = Inches(0.35)
    sec.footer_distance = Inches(0.35)

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    styles["Normal"].font.size = Pt(9.5)
    styles["Normal"].paragraph_format.space_after = Pt(3)
    styles["Normal"].paragraph_format.line_spacing = 1.05

    header = doc.add_table(rows=2, cols=3)
    set_table_width(header, [3300, 3000, 3060])
    for row in header.rows:
        for cell in row.cells:
            set_cell_shading(cell, "FFFFFF")

    cell_text(header.cell(0, 0), BRAND, size=12, bold=True, color=BLUE)
    cell_text(header.cell(0, 1), title, size=12, bold=True, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER)
    cell_text(header.cell(0, 2), "Control No: __________________", size=9, bold=True)
    cell_text(header.cell(1, 0), subtitle, size=8.5, color=INK)
    cell_text(header.cell(1, 1), "Cash / assets must be recorded before issue", size=8.5, color=INK, align=WD_ALIGN_PARAGRAPH.CENTER)
    cell_text(header.cell(1, 2), "Date: ____ / ____ / ______", size=9)
    for row in header.rows:
        for cell in row.cells:
            set_cell_shading(cell, LIGHT_BLUE if row is header.rows[0] else LIGHT_GRAY)

    doc.add_paragraph()
    return doc


def add_section_label(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text.upper())
    style_paragraph(p, size=9.5, bold=True, color=BLUE, after=3)
    return p


def add_field_grid(doc, rows, widths=(1900, 2780, 1900, 2780)):
    table = doc.add_table(rows=len(rows), cols=4)
    set_table_width(table, list(widths))
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            cell_text(table.cell(r, c), text, size=8.8, bold=(c % 2 == 0), color=INK if c % 2 == 0 else "000000")
            if c % 2 == 0:
                set_cell_shading(table.cell(r, c), LIGHT_GRAY)
    return table


def add_signature_grid(doc, labels):
    table = doc.add_table(rows=2, cols=len(labels))
    set_table_width(table, [int(9360 / len(labels))] * len(labels))
    for idx, label in enumerate(labels):
        cell_text(table.cell(0, idx), "\n\n", size=9)
        cell_text(table.cell(1, idx), label + "\nName / Signature / Date", size=8.5, bold=True, color=INK, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(table.cell(1, idx), LIGHT_GRAY)
    return table


def add_footer_note(doc, note):
    p = doc.add_paragraph()
    run = p.add_run(note)
    style_paragraph(p, size=7.8, color="555555", after=0)


def petty_cash_receipt():
    doc = setup_doc("PETTY CASH RECEIPT", "For cash control and expenditure accountability")
    add_section_label(doc, "Transaction Details")
    add_field_grid(doc, [
        ["Paid to / Received by", "", "Department / Project", ""],
        ["Purpose / Description", "", "Account / Cost Code", ""],
        ["Amount in figures", "K / $ ____________", "Amount in words", ""],
        ["Payment type", "Cash / Reimbursement / Advance", "Supporting docs", "Yes / No"],
    ])

    add_section_label(doc, "Cash Control")
    add_field_grid(doc, [
        ["Opening petty cash balance", "", "Amount paid out", ""],
        ["Cash balance after payment", "", "Variance checked", "Yes / No"],
    ])

    add_section_label(doc, "Reason / Notes")
    for _ in range(3):
        add_rule_paragraph(doc, "", 92)

    add_section_label(doc, "Approval and Receipt")
    add_signature_grid(doc, ["Prepared by", "Approved by", "Cashier / Paid by", "Recipient"])
    add_footer_note(doc, "Attach invoices or written authorization where available. No petty cash payment is valid without recipient signature and approval.")
    return doc


def rental_issue_receipt():
    doc = setup_doc("RENTAL ISSUE RECEIPT", "For rentals issued out, deposits, due dates, and return control")
    add_section_label(doc, "Customer / Hirer Details")
    add_field_grid(doc, [
        ["Customer / Tenant name", "", "Phone / Contact", ""],
        ["ID / NRC / Reference", "", "Address / Location", ""],
    ])

    add_section_label(doc, "Rental Details")
    table = doc.add_table(rows=5, cols=6)
    widths = [820, 2700, 900, 1500, 1700, 1740]
    set_table_width(table, widths)
    headers = ["Qty", "Item / Property issued", "Unit", "Period / Due date", "Rate / Deposit", "Condition out"]
    for c, h in enumerate(headers):
        cell_text(table.cell(0, c), h, size=8.3, bold=True, color=INK, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(table.cell(0, c), LIGHT_BLUE)
    for r in range(1, 5):
        for c in range(6):
            cell_text(table.cell(r, c), "", size=8.5)

    add_section_label(doc, "Payment and Return Control")
    add_field_grid(doc, [
        ["Amount paid", "", "Balance due", ""],
        ["Expected return date", "", "Actual return date", ""],
        ["Return condition", "Good / Damaged / Missing", "Damage charge", ""],
    ])

    add_section_label(doc, "Acknowledgement")
    add_signature_grid(doc, ["Issued by", "Authorized by", "Received by customer", "Returned / Checked by"])
    add_footer_note(doc, "Rental items remain company property until formally returned and checked. Record damage, missing parts, or late return charges before closing the rental.")
    return doc


def stock_issue_voucher():
    doc = setup_doc("STOCK ISSUE VOUCHER", "For stock issued from stores to departments, jobs, or staff")
    add_section_label(doc, "Issue Details")
    add_field_grid(doc, [
        ["Issued to", "", "Department / Job", ""],
        ["Requested by", "", "Store / Location", ""],
        ["Reason for issue", "", "Reference / Work order", ""],
    ])

    add_section_label(doc, "Stock Items Issued")
    table = doc.add_table(rows=8, cols=7)
    widths = [720, 1500, 2800, 760, 980, 980, 1620]
    set_table_width(table, widths)
    headers = ["No.", "Stock code", "Description", "Unit", "Qty req.", "Qty issued", "Remarks / batch"]
    for c, h in enumerate(headers):
        cell_text(table.cell(0, c), h, size=8.1, bold=True, color=INK, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(table.cell(0, c), LIGHT_BLUE)
    for r in range(1, 8):
        for c in range(7):
            cell_text(table.cell(r, c), str(r) if c == 0 else "", size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER if c in (0, 3, 4, 5) else WD_ALIGN_PARAGRAPH.LEFT)

    add_section_label(doc, "Inventory Control")
    add_field_grid(doc, [
        ["Posted to stock card / system", "Yes / No", "Posted by", ""],
        ["Remaining stock checked", "Yes / No", "Variance / notes", ""],
    ])

    add_section_label(doc, "Authorization")
    add_signature_grid(doc, ["Requested by", "Approved by", "Issued by stores", "Received by"])
    add_footer_note(doc, "Stock must not leave stores without approval. Quantities issued must be posted to the stock record on the same day.")
    return doc


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    docs = {
        "inventive-harvest-petty-cash-receipt.docx": petty_cash_receipt(),
        "inventive-harvest-rental-issue-receipt.docx": rental_issue_receipt(),
        "inventive-harvest-stock-issue-voucher.docx": stock_issue_voucher(),
    }
    for name, doc in docs.items():
        doc.core_properties.title = name.replace("-", " ").replace(".docx", "").title()
        doc.core_properties.subject = "Inventive Harvest DUB-KIPASA-TC control form"
        doc.core_properties.author = "Inventive Harvest DUB-KIPASA-TC"
        doc.save(OUT_DIR / name)
        print(OUT_DIR / name)


if __name__ == "__main__":
    main()
