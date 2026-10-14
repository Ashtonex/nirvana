"""
Generates professional Microsoft Word (.docx) documents for Inventive Harvest:
1. docs/INVENTIVE_HARVEST_EMPLOYMENT_CONTRACT.docx
2. docs/INVENTIVE_HARVEST_WORKPLACE_POLICIES.docx
"""

from pathlib import Path
from datetime import date
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

DOCS_DIR = Path("docs")
TODAY = date.today().strftime("%d %B %Y")

COLOR_NAVY = "0B2545"
COLOR_PRIMARY = "1F4D78"
COLOR_ACCENT = "2E74B5"
COLOR_MUTED = "50565E"
COLOR_BORDER = "C0C8D0"
COLOR_LIGHT_BG = "F4F6F9"
COLOR_HEADER_BG = "E8EEF5"


def set_cell_shading(cell, fill_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill_hex)


def set_cell_borders(cell, color=COLOR_BORDER, sz="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.find(qn("w:tcBorders"))
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), sz)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def create_base_doc(title, subtitle):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11.0)
    sec.top_margin = Inches(0.75)
    sec.bottom_margin = Inches(0.75)
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.15

    for h_name, sz, col, b, a in [
        ("Heading 1", 14, COLOR_PRIMARY, 10, 4),
        ("Heading 2", 12, COLOR_ACCENT, 8, 3),
        ("Heading 3", 10.5, COLOR_PRIMARY, 6, 2),
    ]:
        h = styles[h_name]
        h.font.name = "Calibri"
        h.font.size = Pt(sz)
        h.font.bold = True
        h.font.color.rgb = RGBColor.from_string(col)
        h.paragraph_format.space_before = Pt(b)
        h.paragraph_format.space_after = Pt(a)

    # Document Header Title
    p_brand = doc.add_paragraph()
    p_brand.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_brand = p_brand.add_run("INVENTIVE HARVEST")
    r_brand.bold = True
    r_brand.font.size = Pt(18)
    r_brand.font.color.rgb = RGBColor.from_string(COLOR_NAVY)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run(title)
    r_title.bold = True
    r_title.font.size = Pt(13)
    r_title.font.color.rgb = RGBColor.from_string(COLOR_PRIMARY)

    if subtitle:
        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_sub = p_sub.add_run(subtitle)
        r_sub.font.size = Pt(9.5)
        r_sub.font.color.rgb = RGBColor.from_string(COLOR_MUTED)

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_meta = p_meta.add_run(f"Operating Locations: Dub Dub | Kipasa | Trade Center | Head Office  •  Date: {TODAY}")
    r_meta.font.size = Pt(8.5)
    r_meta.font.italic = True
    r_meta.font.color.rgb = RGBColor.from_string(COLOR_MUTED)

    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_after = Pt(6)
    r_div = p_div.add_run("―" * 60)
    r_div.font.color.rgb = RGBColor.from_string(COLOR_BORDER)
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER

    return doc


def add_callout_box(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, COLOR_LIGHT_BG)
    set_cell_borders(cell, color=COLOR_ACCENT, sz="8")
    set_cell_margins(cell, top=100, start=140, bottom=100, end=140)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(COLOR_NAVY)
    doc.add_paragraph()


def add_table_data(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Headers
    for c_idx, h_text in enumerate(headers):
        cell = table.cell(0, c_idx)
        cell.text = h_text
        set_cell_shading(cell, COLOR_HEADER_BG)
        set_cell_borders(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.bold = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor.from_string(COLOR_NAVY)

    # Data Rows
    for r_idx, row_data in enumerate(rows):
        for c_idx, val in enumerate(row_data):
            cell = table.cell(r_idx + 1, c_idx)
            cell.text = str(val)
            set_cell_borders(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if (r_idx % 2) == 1:
                set_cell_shading(cell, COLOR_LIGHT_BG)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            for run in p.runs:
                run.font.size = Pt(8.5)

    doc.add_paragraph()
    return table


def build_contract_docx():
    doc = create_base_doc(
        "CONTRACT OF EMPLOYMENT",
        "Confidential & Legally Binding Agreement for Staff at Dub Dub, Kipasa, Trade Center & Head Office"
    )
    add_callout_box(
        doc,
        "NOTICE TO EMPLOYEE AND SUPERVISOR: This is a formal, legally enforceable contract of employment. "
        "Before signing, all particulars including legal names, national identification numbers, remuneration, "
        "assigned duty station, and working hours must be completed accurately."
    )

    doc.add_heading("PARTICULARS OF EMPLOYMENT", level=1)
    particulars = [
        ("Employer", "INVENTIVE HARVEST"),
        ("Operating Network", "Dub Dub Branch, Kipasa Branch, Trade Center Branch, Warehousing & Head Office"),
        ("Employee Full Name", "[Insert Full Legal Name]"),
        ("National ID / NRC / Passport No.", "[Insert NRC / ID / Passport Number]"),
        ("Residential Address", "[Insert Physical Residential Address]"),
        ("Primary Contact Mobile", "[Insert Phone Number(s)]"),
        ("Job Title / Designation", "[Insert: Cashier / Shop Assistant / Branch Supervisor / Store Manager]"),
        ("Assigned Workstation / Branch", "[Insert: Dub Dub / Kipasa / Trade Center / Head Office]"),
        ("Reporting Line", "Branch Supervisor / Store Manager / Managing Director"),
        ("Commencement Date", "[Insert Start Date, e.g., 01 October 2026]"),
        ("Employment Status", "Full-Time (Subject to 3 Months Probation)"),
        ("Gross Monthly Basic Salary", "[Insert Currency & Amount, e.g., USD / ZMW _____]"),
        ("Pay Frequency & Date", "Monthly in Arrears on or before the 28th of each month"),
        ("Standard Working Hours", "44–48 Hours per week across a rotating shift roster (minimum 1 rest day/week)")
    ]
    add_table_data(doc, ["Field / Particular", "Agreed Specification"], particulars)

    clauses = [
        ("1. Appointment & General Duties", [
            "The Employer hereby employs the Employee, and the Employee accepts employment on the terms set forth herein.",
            "The Employee shall perform all duties specified in the Job Description (Schedule A) with diligence, fidelity, and utmost honesty.",
            "Mobility & Transfer: The Employee acknowledges that Inventive Harvest operates multiple branches. The Company reserves the right to rotate or transfer the Employee between Dub Dub, Kipasa, Trade Center, or other branches based on operational needs."
        ]),
        ("2. Probationary Period & Performance Assessment", [
            "The initial period of three (3) months shall constitute probation to assess competence, punctuality, cash control, and customer service.",
            "The Employer may extend probation by up to three (3) additional months where corrective guidance is warranted.",
            "During probation, either Party may terminate employment upon giving one (1) week's written notice or basic salary in lieu thereof."
        ]),
        ("3. Remuneration, Deductions & Overtime", [
            "Remuneration shall be paid into the Employee's designated bank account or registered mobile money wallet on or before the monthly pay date.",
            "Statutory deductions (PAYE, Social Security / Pension, NHIMA) shall be deducted in accordance with applicable laws.",
            "Authorized Recovery: The Employee consents to lawful payroll deductions for unverified cash drawer shortages, unreturned advances, or proven company asset losses occurring under the Employee's custody.",
            "Overtime requires prior written management authorization and shall be compensated in accordance with statutory standards."
        ]),
        ("4. Strict Cash, Till & POS Accountability", [
            "MANDATORY POS RULE: Every customer transaction must be entered into the Point-of-Sale (POS) system (or official pre-numbered manual backup register during downtime) before merchandise leaves the shop.",
            "Every customer must be handed an official printed receipt.",
            "Daily Opening & Closing: Cashiers must perform joint opening and closing float counts with supervisors, signing the daily cash control register.",
            "Strict Bucket Separation: Shop till money, Operations funds, Rent reserve, and Savings / Black Box funds must never be commingled or diverted.",
            "NO PERSONAL BORROWING: Employees are strictly forbidden from taking money from the cash drawer, placing personal IOUs, or keeping personal cash inside the till."
        ]),
        ("5. Stock Control, Inward Receiving & Branch Transfers", [
            "Merchandise received from suppliers or central stores must be counted, checked for quality, and inwarded into the POS before public display.",
            "No merchandise may move between Dub Dub, Kipasa, and Trade Center without an authorized Stock Transfer Voucher (STV) signed by both dispatching and receiving staff.",
            "The Employee must participate fully in weekly cycle counts and monthly physical inventory audits.",
            "Physical Searches: The Employee consents to reasonable physical searches of bags, lockers, and parcels upon entering or exiting company premises."
        ]),
        ("6. Leave & Absence Management", [
            "Paid annual leave accrues at 2 working days per month of service (24 days per annum) and must be requested 14 days in advance.",
            "Sick leave requires notice at least 2 hours before shift start, and a certified medical practitioner certificate for absences over 1 day.",
            "Unapproved absence for three (3) consecutive days without notice constitutes job abandonment and triggers formal contract termination."
        ]),
        ("7. Confidentiality, Non-Compete & Trade Secrets", [
            "The Employee shall perpetually protect all trade secrets, supplier lists, profit margins, turnover data, and POS credentials.",
            "The Employee shall not operate or participate in any competing retail business, sell personal items on company premises, or solicit company clients.",
            "All secret commissions, supplier gifts, or kickbacks are strictly forbidden."
        ]),
        ("8. Disciplinary Code & Termination of Employment", [
            "Disciplinary procedures shall adhere to progressive discipline: Verbal Warning, Written Warning, Final Written Warning, and Summary Dismissal.",
            "Gross Misconduct (leading to immediate dismissal following fair hearing) includes: theft, till manipulation, unrecorded sales, fraudulent records, unauthorized cash borrowing, drunkenness, physical violence, gross insubordination, and breach of confidentiality.",
            "Termination Notice: Following probation, either Party may terminate this Agreement by giving one (1) month's written notice or basic salary in lieu thereof.",
            "Exit Clearance: All keys, passwords, company uniforms, tablets, and property must be surrendered before final terminal dues are released."
        ])
    ]

    for title, points in clauses:
        doc.add_heading(title, level=1)
        for pt in points:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(pt)
            r.font.size = Pt(9.5)

    doc.add_heading("EXECUTION AND SIGNATURES", level=1)
    sig_table = doc.add_table(rows=4, cols=2)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    labels = [
        ("FOR AND ON BEHALF OF INVENTIVE HARVEST", "EMPLOYEE SIGN-OFF"),
        ("Authorized Representative: _______________________", "Employee Full Name: _______________________"),
        ("Signature: ____________________________________", "Signature: ____________________________________"),
        (f"Date: _____ / _____ / 20___   Place: ____________", f"Date: _____ / _____ / 20___   Place: ____________")
    ]
    for r_idx, (col1, col2) in enumerate(labels):
        cell1, cell2 = sig_table.rows[r_idx].cells
        cell1.text = col1
        cell2.text = col2
        set_cell_borders(cell1)
        set_cell_borders(cell2)
        set_cell_margins(cell1)
        set_cell_margins(cell2)
        if r_idx == 0:
            set_cell_shading(cell1, COLOR_HEADER_BG)
            set_cell_shading(cell2, COLOR_HEADER_BG)
            cell1.paragraphs[0].runs[0].font.bold = True
            cell2.paragraphs[0].runs[0].font.bold = True

    out_file = DOCS_DIR / "INVENTIVE_HARVEST_EMPLOYMENT_CONTRACT.docx"
    doc.save(out_file)
    print(f"Saved: {out_file}")


def build_policies_docx():
    doc = create_base_doc(
        "WORKPLACE POLICIES & OPERATING MANUAL",
        "Standard Operating Procedures & Rules for Dub Dub, Kipasa, Trade Center & Operational Hubs"
    )
    add_callout_box(
        doc,
        "CORE MANDATE: Every employee of Inventive Harvest must read, understand, and strictly follow this manual. "
        "It governs all store operations, POS compliance, cash drawer integrity, inventory logistics, and professional standards."
    )

    sections = [
        ("1. Code of Conduct & Business Ethics", [
            "Customer Excellence: Customers must be greeted with warmth and respect. Rude, dismissive, or argumentative conduct is strictly prohibited.",
            "Honesty & Reporting: Employees have a duty to immediately report observed theft, cash discrepancies, unrecorded sales, or misconduct.",
            "Anti-Bribery & Kickbacks: Staff must never accept secret rebates, personal discounts, or favors from suppliers or landlords.",
            "No Moonlighting / Private Sales: Staff are strictly forbidden from selling personal merchandise inside company shops or diverting customers.",
            "Harassment & Bullying: Zero tolerance for sexual harassment, verbal abuse, intimidation, or tribal/gender discrimination."
        ]),
        ("2. Shop Operations & Point-of-Sale (POS) Standards", [
            "Shift Opening: Arrive 15 minutes before opening. Verify locks, inspect premises, conduct joint float counts, and test POS hardware.",
            "The Golden Rule: NO UNRECORDED SALES. Every item leaving the shop must be entered into the POS system before being given to the customer.",
            "Mandatory Official Receipts: Every customer must receive a printed POS receipt. Discarding printed receipts is a disciplinary offense.",
            "Mobile Money & Card Settlements: Verify the exact transaction reference code on the company receiving terminal before releasing goods.",
            "Downtime / Manual Backup: In the event of system or power outage, record all transactions line-by-line in the pre-numbered Manual Sales Register. Inward transactions immediately once system is restored.",
            "Shift Closing: Run the end-of-day Z-report, reconcile physical cash against expected cash, document any variance, and lock up securely."
        ]),
        ("3. Cash Management & Outflow Controls", [
            "Segregation of Buckets: Distinct accounts exist for Shop Till, Operations Fund, Rent Reserve, Savings / Black Box, and Petty Cash. Mixing these funds is strictly prohibited.",
            "No Direct Till Payouts: Expenses must never be paid directly from sales takings. Payouts must come from approved Petty Cash or Operations funds.",
            "Payment Vouchers: All cash outflows must have a numbered payment voucher stating date, expense category, payee, approving supervisor, and attached original receipt.",
            "Approval Thresholds: Supervisor approval up to $20; Managing Director approval required for expenses over $20; expenses over $100 require pre-documentation.",
            "Absolute Ban on IOUs: Borrowing from the till or placing personal chits/notes is treated as criminal misappropriation resulting in dismissal.",
            "Cash Overages: Surplus cash belongs to the Company and must be credited in the daily closing report. Pocketing surplus is theft."
        ]),
        ("4. Inventory Control & Inter-Branch Transfers", [
            "Stock Receiving: Verify incoming stock against delivery invoices for count and condition. Inward goods into the POS system immediately.",
            "Inter-Branch Transfers: Moving goods between Dub Dub, Kipasa, and Trade Center requires a pre-numbered Stock Transfer Voucher (STV) signed by both dispatching and receiving staff.",
            "Stock Audits: Staff must assist with weekly cycle counts, monthly full stocktakes, and unannounced spot checks.",
            "Damaged Goods: Broken or expired merchandise must be quarantined and recorded in the Damaged Goods Register. Staff may not consume or unilaterally discard damaged items."
        ]),
        ("5. Working Hours, Timekeeping & Attendance", [
            "Published Rosters: Roster schedules will be posted 3 days in advance. Informal shift swapping without supervisor approval is prohibited.",
            "Punctuality: Staff must be in uniform on the sales floor ready to trade at start time. Three unexcused late arrivals per month attract disciplinary sanction.",
            "Meal Intervals: Meals must be staggered so counters are never unattended. No eating or open drinks behind the till counter.",
            "Absence Reporting: Notify supervisors of illness at least 2 hours prior to shift. Provide a medical certificate for absences over 1 day."
        ]),
        ("6. Security, Asset Protection & Physical Inspections", [
            "Key Security: Store keys must remain with designated keyholders and must never be duplicated or left on unattended counters.",
            "CCTV Surveillance: Premises and till areas are under video surveillance. Obstructing or tampering with cameras is gross misconduct.",
            "Right of Search: Company reserves the right to conduct reasonable physical searches of bags, lockers, and outer garments upon exit/entry.",
            "Counterfeit Checks: All high-denomination banknotes must be tested with UV detection and security thread checks before acceptance."
        ]),
        ("7. Disciplinary Offenses Matrix", [
            "Minor Offenses (Late arrival, uniform non-compliance, eating at counter): Verbal Warning -> Written Warning -> Final Written Warning.",
            "Serious Offenses (Till shortage, failing to issue receipt, insubordination): Written Warning -> Final Written Warning -> Disciplinary Hearing.",
            "Gross Misconduct (Theft, unrecorded sales, till manipulation, IOUs, intoxication, violence, desertion): Summary Dismissal upon fair hearing."
        ])
    ]

    for title, items in sections:
        doc.add_heading(title, level=1)
        for item in items:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(item)
            r.font.size = Pt(9.5)

    doc.add_heading("EMPLOYEE COMPLIANCE UNDERTAKING & SIGN-OFF", level=1)
    add_callout_box(
        doc,
        "I hereby certify that I have received, read, and understood the Inventive Harvest Workplace Policies & Operating Manual. "
        "I undertake to adhere strictly to all cash handling, POS standards, inventory security, and behavioral rules outlined herein."
    )

    ack_table = doc.add_table(rows=4, cols=2)
    ack_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    ack_fields = [
        ("Employee Full Name: _______________________", "National ID / NRC No.: _______________________"),
        ("Assigned Branch: [ ] Dub Dub  [ ] Kipasa  [ ] Trade Center", "Job Designation: _______________________"),
        ("Employee Signature: _______________________", "Date: _____ / _____ / 20___"),
        ("Supervisor / Witness: _______________________", "Supervisor Signature & Date: _________________")
    ]
    for r_idx, (f1, f2) in enumerate(ack_fields):
        c1, c2 = ack_table.rows[r_idx].cells
        c1.text = f1
        c2.text = f2
        set_cell_borders(c1)
        set_cell_borders(c2)
        set_cell_margins(c1)
        set_cell_margins(c2)
        if r_idx == 0:
            set_cell_shading(c1, COLOR_LIGHT_BG)
            set_cell_shading(c2, COLOR_LIGHT_BG)

    out_file = DOCS_DIR / "INVENTIVE_HARVEST_WORKPLACE_POLICIES.docx"
    doc.save(out_file)
    print(f"Saved: {out_file}")


if __name__ == "__main__":
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    build_contract_docx()
    build_policies_docx()
