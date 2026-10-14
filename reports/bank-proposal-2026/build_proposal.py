from pathlib import Path
import json, calendar, csv
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
src=json.loads((ROOT/'source_extract.json').read_text())
sales=[r for r in src['sales'] if not r.get('deleted_at')]
months=[f'2026-{m:02}' for m in range(3,9)]
def total(shop,month,key): return sum(float(r.get(key) or 0) for r in sales if r['shop_id']==shop and r['date'].startswith(month))
gross=sum(total(s,m,'total_with_tax') for s in ['kipasa','dubdub'] for m in months)
net=sum(total(s,m,'total_before_tax') for s in ['kipasa','dubdub'] for m in months)
base=net/6
money=lambda n:f'{n:,.2f}'
doc=Document()
sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.75)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 st=doc.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.space_after=Pt(7)
doc.styles['Title'].font.size=Pt(25)
doc.styles['Heading 1'].font.size=Pt(16)
doc.styles['Heading 2'].font.size=Pt(12)
for st in doc.styles:
 for el in list(st.element.iter(qn('w:pBdr'))): el.getparent().remove(el)
footer=sec.footer.paragraphs[0];footer.alignment=2
footer.add_run('Inventivee Harvest  |  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
def p(t,style=None):return doc.add_paragraph(t,style)
new_page=False
def h(t):
 global new_page
 heading=doc.add_heading(t,1)
 if new_page:heading.paragraph_format.page_break_before=True;new_page=False
def sub(t):doc.add_heading(t,2)
def page():
 global new_page
 new_page=True
def table(headers,rows,widths=None):
 t=doc.add_table(rows=1, cols=len(headers));t.style='Table Grid';t.autofit=False
 for i,x in enumerate(headers):t.rows[0].cells[i].text=str(x)
 for row in rows:
  c=t.add_row().cells
  for i,x in enumerate(row):c[i].text=str(x)
 for ri,row in enumerate(t.rows):
  trPr=row._tr.get_or_add_trPr();el=OxmlElement('w:cantSplit');trPr.append(el)
  for i,c in enumerate(row.cells):
   if widths:c.width=Inches(widths[i])
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(4);para.paragraph_format.space_before=Pt(4)
    for r in para.runs:r.font.size=Pt(10)
   if ri==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EDF2');c._tc.get_or_add_tcPr().append(sh)
    for r in c.paragraphs[0].runs:r.bold=True
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 return t

doc.add_heading('Inventivee Harvest funding proposal',0)
p('US$10,000 business working capital request','Subtitle')
p('Prepared for BancABC Zimbabwe\n15 September 2026')
h('Funding request')
p('Inventivee Harvest requests a US$10,000 working capital facility to support a government Christmas-decoration tender opportunity in December 2026 and initial stock for a new retail shop intended to begin trading by January 2027. The proposal uses trading records from Kipasa Branch and Dub Dub Outlet to demonstrate existing activity and frame repayment. Nirvana is the financial recording system.')
p('We propose a 12-month facility with separately controlled tender and new-shop stock allocations. The tender allocation would be released against a confirmed award or purchase order, an accepted cost budget and a documented payment schedule. The stock allocation would be paid against supplier quotations and a confirmed opening plan. All terms remain subject to BancABC assessment.')
table(['Proposed facility','Request'],[
 ['Borrower','Inventivee Harvest; exact registered name and legal form to be verified'],
 ['Amount and currency','US$10,000'],
 ['Purpose','Government Christmas decorations and new-shop opening stock'],
 ['Proposed allocation','US$6,000 tender / US$4,000 stock; planning split for management confirmation'],
 ['Timing','Tender delivery in December 2026; shop trading by January 2027'],
 ['Tenor and repayments','12 months; monthly principal and interest; tender receipt sweep if agreed'],
 ['Evidence','Kipasa and Dub Dub records; legal ownership and cash availability to be documented']],[1.7,5.3])
sub('Credit case in brief')
p(f'The two outlets recorded US${money(gross)} in tax-inclusive sales from March to August 2026, averaging US${money(gross/6)} per month. Sales before the tax amounts recorded in the system averaged US${money(base)}. These are transaction records, not audited financial statements or bank-confirmed collections. [S1]')
p('The request is for consideration and structuring. Tender award status, contract value, payment terms, security and the complete borrower cash-flow position are still to be confirmed. The illustrations below deliberately give no repayment credit to tender receipts or new-shop sales.')

page();h('Business and trading evidence')
p('Management identifies Kipasa and Dub Dub as the existing operating shops supporting this request. Recorded items include jewellery, gifts, wrapping paper, balloons and general retail merchandise. Product records support a retail trading profile; they do not establish prior government tender experience. Inventivee Harvest should attach proof that it owns or has lawful access to the cash flows of both outlets. [S1, M1]')
table(['2026 month','Kipasa gross','Dub Dub gross','Combined gross','Before recorded tax'],[[calendar.month_name[int(m[-2:])],money(total('kipasa',m,'total_with_tax')),money(total('dubdub',m,'total_with_tax')),money(sum(total(s,m,'total_with_tax') for s in ['kipasa','dubdub'])),money(sum(total(s,m,'total_before_tax') for s in ['kipasa','dubdub']))] for m in months]+[['Total',money(sum(total('kipasa',m,'total_with_tax') for m in months)),money(sum(total('dubdub',m,'total_with_tax') for m in months)),money(gross),money(net)]],[1.1,1.35,1.35,1.5,1.7])
p('Amounts in US dollars as used in the existing reporting convention; confirm source currency against bank statements. Gross means total_with_tax; before recorded tax means total_before_tax. No statutory tax rate has been inferred. Minor cent differences arise from transaction rounding. [S1]')
sub('Trading continuity and concentration')
p('March through August is the common six-month period with records for both shops. Dub Dub contributes 65.9% of combined gross sales in that period. August combined gross sales were US$13,818.50, 6.9% below June. No December seasonal uplift is assumed. January and February records have substantially fewer entries and require completeness checks before use as a seasonal benchmark. September is incomplete and excluded from the forecast baseline. [S1]')
p('The extraction contains 5,872 sale rows and 2,676 shop-ledger rows for the selected shops within the requested date window. Soft-deleted sales are excluded from the totals. The latest recorded sale is 14 September 2026. Missing periods are not treated as zero trading. Source records alone do not establish 12 months of operating history. [S1]')
sub('Collection evidence')
p('Of the March–August gross sales, US$84,376.84 was marked cash and US$637.00 EcoCash. These payment labels support a predominantly cash retail model, but require reconciliation to daily cash counts, mobile-money statements and bank deposits. Loan monitoring should use actual collections rather than sales alone. [S1]')

page();h('Use of funds and delivery plan')
p('The US$6,000 / US$4,000 allocation below is a proposed planning split, not a supplier quotation or a management-approved budget. Replace it with supported costs before submission. Funding must stay within the US$10,000 approved limit; any extra requirement needs a separately evidenced source.')
table(['Allocation','US$','Purpose and release condition'],[
 ['December tender','6,000','Christmas-decoration procurement and direct fulfilment costs. Release only after award, specifications, quotations, delivery acceptance requirements and payment terms are confirmed.'],
 ['January shop stock','4,000','Initial saleable inventory, including quoted landed costs. Order sufficiently early for January opening; pay suppliers directly where feasible.'],
 ['Total','10,000','Fees, fit-out, deposits and cost overruns require separate funding or a revised allocation.']],[1.4,.8,4.8])
sub('Government Christmas-decoration tender')
p('Management identifies the customer as government. The procuring entity, tender reference, award status, contract price, deposit, invoice due date and payment history have not been supplied. The bank should assess the actual counterparty and contract. A government customer does not make the tender an awarded contract or guarantee prompt payment. [M1]')
p('Before drawing: attach the tender or award document; bill of quantities; supplier quotations; delivery timetable; tender margin after all direct costs, taxes, transport and finance costs; and evidence of how acceptance and invoicing will be completed. Confirm any bid or performance security requirements and whether these need a separate bank instrument.')
p('Proposed control: direct any tender deposit and subsequent receipts to the agreed BancABC collection account. An agreed sweep would repay the tender sub-allocation and accrued charges first, subject to statutory obligations and bank terms. No claim is made that assignment of receivables or a guarantee is already available.')
sub('New retail shop')
p('Opening target: January 2027. Management should confirm the location, lease, product mix, stock quantities, opening date, staffing and launch budget. Separate the US$4,000 initial stock order from rent deposit, fittings, licences and other setup costs. Use demand and sell-through evidence from the existing shops when selecting products. [M1]')
p('The repayment illustration allows US$500 per month for incremental new-shop costs from January, with no new-shop revenue. This is a sensitivity assumption, not an established operating budget. Any additional setup or recurring cost must be included in the final forecast before commitment.')

page();h('Financial interpretation and proposed controls')
sub('Existing profit reports need reconciliation')
old=json.loads((ROOT.parent/'may_june_accurate_financials.json').read_text())
table(['Reported operating result','May 2026','June 2026'],[[shop,money(old['May 2026'][key]['incomeStatement']['operatingProfit']),money(old['June 2026'][key]['incomeStatement']['operatingProfit'])] for shop,key in [('Kipasa','Kipasa Branch'),('Dub Dub','Dubdub Outlet')]], [3,2,2])
p('These are the legacy report outputs, not adjusted earnings or verified debt-service cash flow. The generating script uses tax-inclusive sales and estimates cost of sales from inventory costs; it excludes central Operations activity. Dub Dub shows no separately classified salary expense. Daily till openings must not be accumulated into a bankable cash balance. Reconcile tax treatment, payroll, purchases, central costs, drawings and actual closing cash before using profit or balance-sheet conclusions. [S2]')
p('The July Operations report records US$33,980 in Dub Dub savings and black-box transfers during January–July. It also records pooled Operations uses, including US$17,737 labelled stock orders, and adjustment-looking entries. The report explicitly says these pooled uses cannot all be attributed to individual Dub Dub transfers. The US$33,980 is not added to revenue, profit, security or available cash. [S3]')
sub('Controls proposed for the facility')
p('Maintain separate records for tender spending, new-shop stock and normal retail operations. Match each disbursement to a quotation, payment record and receipt of goods. Reconcile till cash, EcoCash and bank movements daily; review stock and gross margin monthly. Eliminate transfers between the shops and central accounts when preparing consolidated cash flow.')
p('For the loan period, the illustration assumes no owner drawings or discretionary transfers. Any actual drawings must reduce repayment cash in the final forecast. Prioritise statutory payments and loan instalments, and obtain approval before diverting funds to another outlet or purpose. These are proposed controls requiring management adoption, not claims about current practice.')
sub('Security and legal borrower')
p('No collateral or guarantee has been evidenced. Provide the asset and liability statements of the owners, ownership documents, existing encumbrances and details of security proposed to BancABC. Do not represent retail stock or government receivables as accepted security until the bank has assessed them. Document the relationship between Inventivee Harvest and the two trading shops.')

page();h('Twelve month repayment illustration')
p('This is a conditional planning case, not a forecast certified by management or a BancABC quotation. Period: December 2026 to November 2027. It assumes US$10,000 is advanced in December, US$6,000 tender spending occurs in December and US$4,000 opening stock is purchased in January. Earlier procurement requires earlier approval and a revised schedule.')
table(['Input','Illustrative assumption'],[
 ['Existing shop collections',f'US${money(base)} monthly before recorded sales tax; March–August average, no growth'],
 ['Stock replenishment','50% of sales before recorded tax, paid in the same month'],
 ['Existing shop operating costs','US$4,000 monthly; planning allowance, not verified total spending'],
 ['Other cash reserve','US$300 monthly for cash tax and maintenance; not a calculated tax liability'],
 ['New shop incremental costs','US$500 monthly from January; zero new-shop sales included'],
 ['Other flows','No tender receipts, owner drawings or existing debt service included; confirm actual obligations'],
 ['Loan terms','24% nominal annual interest, monthly reducing balance; 12 equal principal instalments'],
 ['Opening unrestricted cash','US$0; no unverified savings or till balance counted']],[2.05,4.95])
p('Revenue is assumed to be collected in the same month. The 50% stock allowance is a cash replenishment assumption; it is not verified historical cost of sales. Receivables, supplier credit, VAT settlement differences and inventory timing need separate reconciliation. Loan fees and insurance are excluded because no quote is available. [A1]')
rows=[];bal=10000;cash=0;csvrows=[]
for i in range(12):
 month='Dec 2026' if i==0 else f'{calendar.month_abbr[i]} 2027'
 op=base*.5-4000-300-(500 if i>=1 else 0)
 interest=bal*.24/12;principal=10000/12;service=principal+interest
 draw=10000 if i==0 else 0;use=6000 if i==0 else 4000 if i==1 else 0
 cash+=draw+op-use-service;bal-=principal
 rows.append([month,money(op),money(draw),money(use),money(service),money(cash)])
 csvrows.append(dict(month=month,collections=base,replenishment=base*.5,existing_costs=4000,reserve=300,new_shop_costs=500 if i else 0,operating_cash=op,loan_draw=draw,project_spend=use,principal=principal,interest=interest,debt_service=service,closing_cash=cash,closing_debt=max(0,bal)))
table(['Month','Operating cash','Loan draw','Project spend','Debt service','Closing cash'],rows,[1.1,1.3,1.05,1.15,1.2,1.2])
p('Operating cash is after the above replenishment, costs and reserve, before new project spending and loan service. Closing cash includes loan proceeds and is not profit. Unspent project allocations must remain reserved. Monthly calculations use unrounded values. [A1]')

page();h('Repayment capacity and downside')
post=base*.5-4800
table(['Case after January opening','Cash before debt','Coverage of US$1,033 first-payment benchmark'],[
 ['Base assumptions',money(post),f'{post/(10000/12+200):.2f}x'],
 ['Sales 15% lower at 50% margin',money(base*.85*.5-4800),f'{(base*.85*.5-4800)/(10000/12+200):.2f}x'],
 ['Margin 45% at baseline sales',money(base*.45-4800),f'{(base*.45-4800)/(10000/12+200):.2f}x'],
 ['Sales down 15% and margin 45%',money(base*.85*.45-4800),f'{(base*.85*.45-4800)/(10000/12+200):.2f}x']],[3,1.3,2.7])
p('Coverage here means modelled cash before debt divided by the first monthly loan payment, used consistently as a conservative comparison. It is not a historical debt-service coverage ratio or BancABC covenant. The base case depends on the cost limits, zero drawings and absence of other debt. The downside cases cannot support the same instalments without additional verified cash or a revised structure. [A1]')
sub('Price sensitivity and repayment boundary')
table(['Assumed annual rate','First payment','Total interest over 12 months'],[[f'{r}%',money(10000/12+10000*r/1200),money(10000*r/100*6.5/12)] for r in [18,24,30]],[2,2.5,2.5])
p(f'At the base 24% illustrative rate, total interest is US$1,300 before fees. With US$4,800 monthly costs and reserve after opening and a 50% cash margin, existing-shop collections before recorded sales tax must be at least US${money((4800+(10000/12+200))/.5)} to meet the first-payment benchmark. A self-imposed 1.25x planning buffer requires US${money((4800+1.25*(10000/12+200))/.5)} monthly collections. The baseline of US${money(base)} is only modestly above that buffer. These thresholds are planning tests, not bank criteria. [A1]')
sub('Recommended structure and fallback')
p('Request one US$10,000 working capital approval with two monitored allocations. Discuss staged releases to reduce interest on unused funds, subject to bank availability. The displayed schedule conservatively charges interest on a full December draw. Agree whether tender receipts permit or require early repayment, and obtain the applicable charge in writing.')
p('If the tender is not awarded, do not draw its allocation. If the verified cost budget or liabilities fail the repayment test, revise the stock rollout, evidence an owner contribution or discuss a smaller eligible facility with the bank. The published business-loan minimum is US$10,000, so a smaller request may need a different product. A longer tenor is an alternative to discuss, not an assumed working capital entitlement. [B1, B2]')

page();h('Submission requirements and next steps')
p('BancABC publishes business-loan requirements covering operating history, cash flow, financial records, security and an established account relationship. Its working capital offering describes facilities up to 12 months. Product eligibility and final conditions should be confirmed directly with the bank. [B1, B2]')
table(['Priority','Item to complete','Evidence required'],[
 ['1','Borrower identity and bank relationship','Registration, KYC, owners and signatories; proof shops belong to or support borrower; BancABC account tenure.'],
 ['1','Tender facts and economics','Government entity, tender reference, award or order, contract price, quotations, delivery acceptance and payment schedule.'],
 ['1','Final US$10,000 allocation','Confirm tender/stock split; add finance fees, bid/performance security, setup costs and source of any shortfall.'],
 ['1','Cash flow and financial reconciliation','Bank and mobile-money statements; latest management accounts; stock counts and purchase costs; tax, payroll, drawings, debt and central Operations reconciliation.'],
 ['1','Security and liabilities','Owners’ assets and liabilities; existing debt terms; title and value of proposed security, subject to bank assessment.'],
 ['2','New shop execution','Location, lease, January launch schedule, stock list, staffing and actual operating budget.'],
 ['2','Trading history and approvals','Evidence of at least 12 months operating history; financial history requested by bank; authority to borrow and signed application.']],[.65,2,4.35])
p('Published business-loan criteria include more than 12 months of operation, more than six months as a BancABC account holder, a minimum US$10,000 request, financial information no older than 12 months, a business plan and at least a 12-month cash-flow forecast. The bank asks for financial history, ordinarily three years, and owners’ asset/liability and security details. Confirm eligibility or any exception before treating the timetable as achievable. [B1]')
sub('Proposed timetable')
p('September–October 2026: complete legal, financial and tender evidence; obtain supplier quotations and bank terms. November: secure facility approval and agree draw conditions before placing funded orders. December: execute awarded tender and procure shop stock to meet supplier lead times. January 2027: begin new-shop trading and monthly bank reporting. The December draw in the model must be moved earlier if procurement requires it.')
sub('Applicant confirmation')
p('Authorised representative ____________________\nCapacity ____________________   Date ____________________\nContact details ____________________')
p('Confirm the financial schedules, final allocation, tender status, new-shop costs and proposed operating controls before signing and submitting this proposal.')

page();h('Sources and calculation notes')
for title,txt in [
 ('M1 Management instructions','Conversation dated 15 September 2026: borrower name Inventivee Harvest; US$10,000 request; government Christmas-decoration tender in December; new-shop stock and January 2027 opening. Award status and contract terms not provided.'),
 ('S1 Nirvana transaction extract',f'Read-only extraction at {src["extractedAt"]}. Tables: shops, sales and ledger_entries; shop IDs kipasa and dubdub; dates from 1 September 2025 to before 16 September 2026. Monthly analysis groups transaction date strings and excludes deleted_at records. Supporting file: source_extract.json; analysis uses March–August 2026. Date-boundary and cash reconciliation should be confirmed by the accountant.'),
 ('S2 Existing financial reports','reports/may_june_accurate_financials.json and scripts/generate_may_june_accurate_financials.js. These legacy POS-only reports are not adopted as audited or complete financial statements. The new sales extract ties to the four May–June gross sales figures to rounding precision.'),
 ('S3 Operations reconciliation','reports/dubdub_operations_use_of_funds_jan_to_jul_31_2026.json, generated 31 July 2026. Pooled Operations uses are not uniquely attributable to Dub Dub. No transfer total is treated as new income or unencumbered cash.'),
 ('A1 Planning model','Analyst assumptions, not lender terms or management-approved forecasts. Existing net sales = US$73,612.03 / 6 = US$12,268.6717 monthly. Cash margin 50%; existing costs US$4,000; tax/maintenance reserve US$300; new-shop costs US$500 from January; owner drawings zero. Principal = US$10,000 / 12; interest = opening monthly principal balance × annual rate / 12. Closing cash = opening cash + collections − replenishment − costs − reserve + draw − project spend − principal − interest. No tender receipts or new-shop sales are included.'),
 ('B1 BancABC business borrowing','Official product page accessed 15 September 2026: https://www.bancabc.co.zw/business-borrow/'),
 ('B2 BancABC working capital','Official product page accessed 15 September 2026: https://www.bancabc.co.zw/business-working-capital-loans/')]:
 sub(title);p(txt)
p('No interest rate, approval probability, credit limit, collateral acceptance or government payment assurance has been represented as a bank commitment. This proposal remains subject to completion of the submission evidence and management confirmation.')
out=ROOT/'Inventivee_Harvest_BancABC_Funding_Proposal.docx'
doc.save(out)
with (ROOT/'cashflow_support.csv').open('w',newline='') as f:
 w=csv.DictWriter(f,fieldnames=list(csvrows[0]));w.writeheader();w.writerows(csvrows)
(ROOT/'calculation_summary.json').write_text(json.dumps({'gross_sales':gross,'net_sales':net,'baseline_monthly':base,'cashflow':csvrows},indent=2))
manifest={'manifest_version':'1.0','skill':'capital-markets-issuance','artifact_mode':'native_document','output_dir':str(ROOT),'first_read':{'path':str(out),'role':'primary human deliverable','why':'Editable lender proposal with financial evidence and repayment illustration'},'primary_human_deliverable':str(out),'human_deliverables':[{'path':str(out),'role':'human_deliverable','artifact_type':'native_document','description':'Funding proposal','user_visible_default':True,'contains_new_analysis':True}],'companion_deliverables':[],'support_artifacts':[{'path':x,'role':'support_artifact','artifact_type':x.split('.')[-1],'description':'Financial evidence or calculation backup','user_visible_default':False,'contains_new_analysis':x!='source_extract.json','support_reason':'Audit trail and reproducibility','user_requested_machine_readable':False} for x in ['source_extract.json','cashflow_support.csv','calculation_summary.json']],'agent_artifacts':[],'support_artifacts_user_visible_default':False,'blocked_or_partial_status':{'status':'partial','reason':'Proposal drafted; external submission requires factual completion','missing_inputs':['Tender award and payment terms','Confirmed uses split','Borrower ownership and eligibility','Security and liabilities','Reconciled financials and actual shop costs']},'final_response_guidance':{'lead_with':'primary_human_deliverable','mention_support_artifacts':'only_briefly_unless_requested'},'discipline_note':'Support artifacts are for audit only.'}
(ROOT/'manifest.json').write_text(json.dumps(manifest,indent=2))
print(out)
print('base',base,'post_opening_cash',post,'end_cash',cash)
